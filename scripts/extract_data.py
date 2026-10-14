"""
Extracts employee / establishment data from the UoK source files into data/seed.json.

Sources (project root):
  - Updated Budget.xlsx                      -> units (departments & offices), budgeted posts, non-teaching staff
  - List of Teachers 12-10-2026.xlsx         -> faculty members (teaching staff) incl. administrative charges
  - Vacant Positions BPS-1  to BPS-16.pdf    -> non-teaching vacancy statement (BPS 1-16)
  - Vacant Positions BPS-17 & above.pdf      -> non-teaching vacancy statement (BPS 17+)

Usage:  python scripts/extract_data.py
"""
import datetime as dt
import difflib
import json
import re
from pathlib import Path

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "seed.json"


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = re.sub(r"\s+", " ", v.replace("�", "'")).strip()
        return v or None
    return v


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str):
        for fmt in ("%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def to_int(v):
    if v is None:
        return None
    m = re.search(r"\d+", str(v))
    return int(m.group()) if m else None


def classify_unit(name):
    n = name.upper()
    if n.startswith("DEAN"):
        return "Faculty / Dean Office"
    if n.startswith("DEPARTMENT") or n.startswith("SCHOOL"):
        return "Teaching Department"
    if n.startswith("INSTITUTE"):
        return "Institute"
    if "CENTRE" in n or "CENTER" in n or "CHAIR" in n or "HERBARIUM" in n:
        return "Research Centre"
    if n.startswith("DIRECTORATE"):
        return "Directorate"
    return "Administrative Office"


def title_case(s):
    small = {"of", "and", "&", "for", "the", "in"}
    words = s.lower().split()
    out = []
    for i, w in enumerate(words):
        if i and w in small:
            out.append(w)
        elif re.fullmatch(r"\(?[a-z]{1,4}\)?", w) and w.strip("()") in {"ubit", "kuh", "rfc", "lan", "asrb", "mphil."}:
            out.append(w.upper())
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


ACRONYMS = {"IT", "PA", "PS", "AV", "LAN", "II", "III", "IV", "QEC", "ACP", "CCT", "UDC", "LDC", "DDO", "NTN", "VC", "PVC"}


def nice_designation(d):
    """'JNR.ACCOUNTANT' -> 'Jnr. Accountant', 'U. D. C.' -> 'U.D.C.', 'SUB - ENGINEER' -> 'Sub-Engineer'"""
    if not d:
        return d
    d = re.sub(r"\s*-\s*", "-", d)
    d = re.sub(r"\.(?=[A-Za-z]{2})", ". ", d)
    d = re.sub(r"(?<=\b[A-Za-z]\.)\s+(?=[A-Za-z]\.)", "", d)
    d = re.sub(r"(?<=\b[A-Za-z]\.)\s+(?=[A-Za-z]\.)", "", d)

    def word(w):
        if re.fullmatch(r"(?:[A-Za-z]\.){2,}", w) or w.strip("().").upper() in ACRONYMS:
            return w.upper()
        return "-".join(p[:1].upper() + p[1:].lower() for p in w.split("-"))

    return " ".join(word(w) for w in d.split())


def guess_gender(name):
    n = (name or "").lower()
    if re.match(r"^(mrs|ms|miss)\b", n) or re.search(r"\b(begum|bano|khatoon|bibi)\b", n):
        return "Female"
    if re.match(r"^mr\b", n) or re.search(r"\b(muhammad|mohammad|ahmed|ahmad|s/o)\b", n):
        return "Male"
    return None


# ---------------------------------------------------------------- budget book
def parse_budget():
    wb = openpyxl.load_workbook(ROOT / "Updated Budget.xlsx", data_only=True)
    units, posts, staff = [], [], []
    for ws in wb.worksheets:
        rows = [tuple(clean(c) for c in r) for r in ws.iter_rows(values_only=True)]
        unit = None
        mode = None
        section = None
        for i, r in enumerate(rows):
            a = r[0]
            nxt = rows[i + 1][0] if i + 1 < len(rows) else None
            if isinstance(a, str) and isinstance(nxt, str) and nxt.startswith("S.No:"):
                code = to_int(nxt)
                unit = {
                    "code": f"{code:03d}",
                    "name": title_case(a),
                    "source_name": a,
                    "sheet": ws.title,
                    "category": classify_unit(a),
                }
                units.append(unit)
                mode, section = None, None
                continue
            if unit is None:
                continue
            if a == "S.No" and r[1] == "Working Designation":
                mode = "budget"
                continue
            if isinstance(a, str) and a.startswith("Section"):
                section = a.replace("Section", "").strip(" '\"“”‘’")
                continue
            if a == "PRESENT STAFF POSITION":
                mode = "staff"
                continue
            if not isinstance(a, (int, float)) or not isinstance(r[1], str):
                continue
            if mode == "budget":
                posts.append({
                    "unit_code": unit["code"],
                    "section": section,
                    "designation": r[1],
                    "bps": to_int(r[2]),
                    "positions": to_int(r[3]) or 0,
                })
            elif mode == "staff":
                staff.append({
                    "unit_code": unit["code"],
                    "name": r[1].upper(),
                    "bps": to_int(r[2]),
                    "designation": r[3],
                    "posting_label": r[4],
                    "date_of_retirement": to_date(r[6]),
                })
    return units, posts, staff


# ---------------------------------------------------------------- faculty list
def parse_teachers():
    wb = openpyxl.load_workbook(ROOT / "List of Teachers 12-10-2026.xlsx", data_only=True)
    ws = wb["Sheet1"]
    out = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        raw = list(r)
        r = [clean(c) for c in raw]
        if not r[2]:
            continue
        status_raw = r[1] or "Working"
        charges = []
        if r[7]:
            charges.append({
                "role": r[7].strip("() "),
                "unit": r[4],
                "present_from": to_date(raw[8]), "present_to": clean(str(raw[9])) if raw[9] and not to_date(raw[9]) else to_date(raw[9]),
                "second_term": pair_terms(raw[10], raw[11]),
                "first_from": to_date(raw[12]), "first_to": to_date(raw[13]),
            })
        if r[15]:
            charges.append({
                "role": r[16] or "Additional Charge",
                "unit": r[15],
                "present_from": to_date(raw[17]),
                "present_to": to_date(raw[18]) or clean(str(raw[18])) if raw[18] else None,
                "second_term": None, "first_from": None, "first_to": None,
            })
        out.append({
            "name": r[2],
            "status_raw": status_raw,
            "designation": r[3],
            "department": r[4],
            "date_of_birth": to_date(raw[5]),
            "date_of_retirement": to_date(raw[6]),
            "director_id": r[14],
            "charges": charges,
        })
    return out


def pair_terms(frm, to):
    """Cells may stack several terms on separate lines: pair the i-th 'from' with the i-th 'to'."""
    split = lambda v: [x.strip() for x in str(v).splitlines() if x.strip()] if v else []
    fmt = lambda v: to_date(v) and dt.date.fromisoformat(to_date(v)).strftime("%d %b %Y") or (v.isoformat()[:10] if hasattr(v, "isoformat") else v)
    f, t = split(frm) if not hasattr(frm, "isoformat") else [frm], split(to) if not hasattr(to, "isoformat") else [to]
    pairs = [f"{fmt(f[i]) if i < len(f) else '?'} to {fmt(t[i]) if i < len(t) else '?'}" for i in range(max(len(f), len(t)))]
    return "; ".join(pairs) or None


def map_status(s):
    s = (s or "").upper()
    if "SUSPEND" in s:
        return "Suspended"
    if "LIEN" in s:
        return "On Lien"
    if "F.S.L" in s:
        return "Foreign Study Leave"
    if "S/L" in s:
        return "Study Leave"
    if "UA" in s or "U/A" in s:
        return "Unauthorized Absence"
    return "Working"


# ---------------------------------------------------------------- vacancies
def parse_vacancies():
    out = []
    files = [
        ("Vacant Positions BPS-1  to BPS-16.pdf", "BPS 1-16", False),
        ("Vacant Positions BPS-17 & above.pdf", "BPS 17 & above", True),
    ]
    for fname, group, has_dept in files:
        with pdfplumber.open(ROOT / fname) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    for row in table:
                        row = [clean(c) for c in row]
                        if not row or not row[0] or not str(row[0]).isdigit():
                            continue
                        if has_dept:
                            _, post, dept, bps, sanc, work, vac, prio = (row + [None] * 8)[:8]
                        else:
                            _, post, bps, sanc, work, vac, prio = (row + [None] * 7)[:7]
                            dept = None
                        out.append({
                            "group": group, "post": post, "department": dept, "bps": bps,
                            "sanctioned": to_int(sanc) or 0, "working": to_int(work) or 0,
                            "vacant": to_int(vac) or 0, "priority": to_int(prio),
                        })
    return out


def norm(s):
    s = (s or "").lower()
    s = re.sub(r"^(department|institute|centre|center|dean,?|school) (of|for)?\s*", "", s)
    s = s.replace("&", "and").replace("agribus.", "agribusiness").replace("phy.", "physical").replace("chem.", "chemical")
    return re.sub(r"[^a-z ]", "", s).strip()


def main():
    units, posts, staff = parse_budget()
    teachers = parse_teachers()
    vacancies = parse_vacancies()

    # map teacher departments onto budget units; create teaching units that do not exist in budget book
    by_norm = {norm(u["source_name"]): u for u in units}
    next_code = 900
    dept_map = {}
    aliases = {
        "National Nematological Research Centre": "NATIONAL NEMATOLOGICAL CENTRE",
        "QEC": "QUALITY ENHANCEMENT CELL",
        "BCC&T": "BUREAU OF COMPOSITION, COMPILATION & TRANCELATION",
        "Registrar": "REGISTRAR OFFICE",
        "Evening Program": "DIRECTORATE OF EVENING PROGRAM",
        "Controller of Examinations": "EXAMINATION OFFICE",
        "Student Guidance, Placement and Overseas Examination Bureau": "STUDENTS GUIDANCE COUNSELLING & PLACEMENT BUREAU",
        "Dr. Muhammad Ajmal Khan Institute of Sustainable Halophyte Utilization": "INSTITUTE OF SUSTAINABLE HALOPHYTE",
        "Dr. Zafar H. Zaidi National Centre for Proteomics": "NATIONAL CENTRE FOR PROTEOMICS",
        "Inst. of Envirnmental Studies": "INSTITUTE OF ENVIRONMENTAL STUDIES",
    }
    for dname in sorted({t["department"] for t in teachers if t["department"]}):
        key = norm(aliases.get(dname, dname))
        match = by_norm.get(key)
        if not match:
            cand = difflib.get_close_matches(key, list(by_norm), n=1, cutoff=0.82)
            match = by_norm[cand[0]] if cand else None
        if not match:
            next_code += 1
            match = {"code": f"{next_code}", "name": dname, "source_name": dname,
                     "sheet": "Faculty List", "category": "Teaching Department"}
            units.append(match)
            by_norm[key] = match
        dept_map[dname] = match["code"]

    def match_unit(name):
        key = norm(aliases.get(name, name))
        if key in by_norm:
            return by_norm[key]["code"]
        cand = difflib.get_close_matches(key, list(by_norm), n=1, cutoff=0.75)
        return by_norm[cand[0]]["code"] if cand else None

    for t in teachers:
        for i, c in enumerate(t["charges"]):
            # first charge is headship of own department; additional charges name another unit
            c["unit_code"] = dept_map.get(t["department"]) if i == 0 and c["unit"] == t["department"] else match_unit(c["unit"] or "")

    employees = []
    for i, t in enumerate(teachers, 1):
        employees.append({
            "employee_no": f"UOK-T-{i:04d}",
            "full_name": t["name"],
            "gender": guess_gender(t["name"]),
            "staff_type": "Teaching",
            "designation": t["designation"],
            "bps": {"Professor": 21, "Associate Professor": 20, "Assistant Professor": 19, "Lecturer": 18}.get(t["designation"]),
            "unit_code": dept_map.get(t["department"]),
            "status": map_status(t["status_raw"]),
            "status_remarks": t["status_raw"] if t["status_raw"] != "Working" else None,
            "date_of_birth": t["date_of_birth"],
            "date_of_retirement": t["date_of_retirement"],
            "director_id": t["director_id"],
            "charges": t["charges"],
            "source": "List of Teachers 12-10-2026.xlsx",
        })
    for i, s in enumerate(staff, 1):
        dob = None
        if s["date_of_retirement"]:
            d = dt.date.fromisoformat(s["date_of_retirement"]) + dt.timedelta(days=1)
            try:
                dob = d.replace(year=d.year - 60).isoformat()
            except ValueError:
                dob = d.replace(year=d.year - 60, day=28).isoformat()
        employees.append({
            "employee_no": f"UOK-N-{i:05d}",
            "full_name": s["name"].title().replace("S/O", "s/o").replace("D/O", "d/o"),
            "gender": guess_gender(s["name"]),
            "staff_type": "Non-Teaching",
            "designation": nice_designation(s["designation"]),
            "bps": s["bps"],
            "unit_code": s["unit_code"],
            "status": "Working",
            "status_remarks": None,
            "date_of_birth": dob,
            "dob_derived": dob is not None,
            "date_of_retirement": s["date_of_retirement"],
            "posting_label": s["posting_label"],
            "charges": [],
            "source": "Updated Budget.xlsx",
        })

    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps({
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "units": units, "budget_posts": posts, "employees": employees, "vacancies": vacancies,
    }, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"units={len(units)} posts={len(posts)} teachers={len(teachers)} non_teaching={len(staff)} vacancies={len(vacancies)}")
    print("teacher depts added outside budget book:", [u["name"] for u in units if u["sheet"] == "Faculty List"])


if __name__ == "__main__":
    main()
